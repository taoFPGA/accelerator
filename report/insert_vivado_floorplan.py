"""
One-shot, targeted edit: replaces the Section V.C "[Figure Placeholder:
Vivado Device View / Floorplan ...]" callout box in the CURRENT
taoFPGA_Project_Book.docx with a real screenshot, and renumbers every
subsequent figure (old Fig. 5-8 -> new Fig. 6-9) throughout the document
(captions, in-text references, and the List of Figures) to make room for
the new Fig. 5.

Operates directly on the live docx in place -- does NOT go through
generate_project_book.py, which would discard the user's manual edits.

Run once from the dev machine:
    python insert_vivado_floorplan.py
"""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

PATH = "taoFPGA_Project_Book.docx"
IMAGE_PATH = r"C:\Users\elira\OneDrive\Pictures\Screenshots\Screenshot 2026-08-28 011709.png"
NEW_CAPTION = (
    "Fig. 5. Vivado Device view (post-implementation): physical cell "
    "placement across the Zynq-7020 fabric, with column-structured "
    "DSP/BRAM resources (yellow) and clock-region boundaries (X0Y0\u2013X1Y2)."
)


def find_placeholder_table(doc):
    for table in doc.tables:
        text = table.rows[0].cells[0].text
        if text.strip().startswith("[Figure Placeholder: Vivado Device View"):
            return table
    raise RuntimeError("Vivado floorplan placeholder table not found")


def insert_figure_before(doc, anchor_element):
    """Builds a centered image paragraph + centered italic caption
    paragraph (matching every other figure() call's exact style: Pt(9)
    italic Times New Roman caption), and splices both in immediately
    before anchor_element in the document body."""
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(IMAGE_PATH, width=Inches(3.4))

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap_p.add_run(NEW_CAPTION)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r.font.italic = True

    # Both paragraphs were appended at the end of the body by add_paragraph();
    # move them into place immediately before the placeholder table.
    anchor_element.addprevious(img_p._p)
    anchor_element.addprevious(cap_p._p)


def renumber(doc):
    """Old Fig. 5/6/7/8 -> new Fig. 6/7/8/9, everywhere they appear as a
    caption or an in-text reference. Walked newest-number-first so a
    caption text search for 'Fig. 7' doesn't accidentally match a
    just-renumbered 'Fig. 7' that used to be 'Fig. 6'."""
    replacements = [
        ("Fig. 8. Real ViT self-attention and MLP dataflows",
         "Fig. 9. Real ViT self-attention and MLP dataflows"),
        ("Fig. 8. Real ViT dataflows contrasted",
         "Fig. 9. Real ViT dataflows contrasted"),
        ("Fig. 7. Measured hardware kernel speedup",
         "Fig. 8. Measured hardware kernel speedup"),
        ("Fig. 6. Throughput (GOP/s)",
         "Fig. 7. Throughput (GOP/s)"),
        ("Fig. 5. Latency, CPU vs. hardware",
         "Fig. 6. Latency, CPU vs. hardware"),
        ("visualized in Figs. 5\u20137", "visualized in Figs. 6\u20138"),
    ]
    for p in doc.paragraphs:
        text = p.text
        new_text = text
        for old, new in replacements:
            if old in new_text:
                new_text = new_text.replace(old, new)
        if new_text != text:
            # Replace only within the first run that actually contains the
            # changed substring is unsafe if the text is split across
            # multiple runs (confirmed pattern in this file after manual
            # edits -- Word often splits runs). Rebuild from the full
            # paragraph text instead, preserving the first run's
            # formatting for the whole rebuilt line, since every one of
            # these caption/body lines is uniformly styled already.
            if not p.runs:
                continue
            font_name = p.runs[0].font.name
            font_size = p.runs[0].font.size
            italic = p.runs[0].font.italic
            bold = p.runs[0].font.bold
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            run = p.add_run(new_text)
            run.font.name = font_name
            if font_size is not None:
                run.font.size = font_size
            run.font.italic = italic
            run.font.bold = bold


def insert_lof_entry(doc):
    """Inserts the new 'Fig. 5. ...' line into the List of Figures,
    immediately after the existing Fig. 4 entry, matching that entry's
    exact tab-stop/leader paragraph formatting (copied wholesale, only the
    run text differs)."""
    fig4_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith("Fig. 4. Complete top-level SoC architecture"):
            fig4_idx = i
            break
    if fig4_idx is None:
        raise RuntimeError("Could not find the Fig. 4 List-of-Figures entry")
    fig4_p = doc.paragraphs[fig4_idx]

    # New page number: this is a genuinely new figure inserted into the
    # body, so the real page can only be read off a rendered PDF -- use a
    # placeholder here and fix it after the next render, exactly as during
    # initial construction of this document.
    label = "Fig. 5. Vivado Device view: physical placement and DSP/BRAM column structure."
    page_placeholder = "#"

    # Simplest reliable approach: clone Fig.4's paragraph XML (preserves
    # its tab-stop definition exactly), insert the clone after it, then
    # overwrite the clone's run text.
    import copy
    clone = copy.deepcopy(fig4_p._p)
    fig4_p._p.addnext(clone)
    from docx.text.paragraph import Paragraph
    clone_p = Paragraph(clone, fig4_p._parent)
    for r in list(clone_p.runs):
        r._element.getparent().remove(r._element)
    run = clone_p.add_run(f"{label}\t{page_placeholder}")
    src_run = fig4_p.runs[0]
    run.font.name = src_run.font.name
    run.font.size = src_run.font.size
    run.font.italic = src_run.font.italic
    run.font.bold = src_run.font.bold


def main():
    doc = docx.Document(PATH)
    placeholder_table = find_placeholder_table(doc)
    anchor = placeholder_table._element
    insert_figure_before(doc, anchor)
    anchor.getparent().remove(anchor)
    renumber(doc)
    insert_lof_entry(doc)
    doc.save(PATH)
    print("Inserted Fig. 5, renumbered Figs. 6-9, updated List of Figures.")


if __name__ == "__main__":
    main()
