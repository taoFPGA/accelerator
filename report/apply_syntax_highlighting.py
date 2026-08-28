"""
One-shot, targeted edit: adds Verilog syntax coloring to the two code
listing tables (right_shifter.v excerpt, PE.v excerpt) in the CURRENT
taoFPGA_Project_Book.docx, in place.

Deliberately does NOT go through generate_project_book.py / the
Document()-from-scratch pipeline -- the user has since made manual
formatting edits directly in Word, and regenerating from the generator
scripts would discard those. This script opens the live docx, touches
only the two listing tables (identified by content, not by assuming a
fixed table index), and saves back to the same path.

Run once from the dev machine:
    python apply_syntax_highlighting.py
"""
import re

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

PATH = "taoFPGA_Project_Book.docx"

KEYWORDS = {
    "module", "endmodule", "input", "output", "inout", "reg", "wire",
    "wand", "wor", "always", "always_ff", "always_comb", "begin", "end",
    "if", "else", "case", "casex", "casez", "endcase", "default", "assign",
    "parameter", "localparam", "generate", "endgenerate", "genvar", "for",
    "while", "posedge", "negedge", "signed", "unsigned", "integer", "real",
    "function", "endfunction", "task", "endtask", "initial",
}

COLOR_COMMENT = RGBColor(0x00, 0x80, 0x00)
COLOR_KEYWORD = RGBColor(0x00, 0x00, 0xEE)
COLOR_STRING = RGBColor(0xA3, 0x15, 0x15)
COLOR_NUMBER = RGBColor(0x09, 0x86, 0x58)

COLOR_MAP = {
    "comment": COLOR_COMMENT,
    "string": COLOR_STRING,
    "vlit": COLOR_NUMBER,
    "num": COLOR_NUMBER,
    "keyword": COLOR_KEYWORD,
}

TOKEN_RE = re.compile(
    r"(?P<comment>//.*$)"
    r"|(?P<string>\"[^\"]*\")"
    r"|(?P<vlit>\d+'[bBoOdDhH][0-9a-fA-Fxz_]+)"
    r"|(?P<num>\b\d+\b)"
    r"|(?P<ident>[A-Za-z_][A-Za-z0-9_]*)"
    r"|(?P<ws>\s+)"
    r"|(?P<punct>[^\w\s]+)"
)


def tokenize(line):
    for m in TOKEN_RE.finditer(line):
        kind = m.lastgroup
        text = m.group()
        if kind == "ident" and text in KEYWORDS:
            kind = "keyword"
        yield kind, text


def colorize_paragraph(p):
    text = p.text
    if not text:
        return
    font_name = p.runs[0].font.name if p.runs else "Consolas"
    font_size = p.runs[0].font.size if p.runs else None
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for kind, tok in tokenize(text):
        run = p.add_run(tok)
        run.font.name = font_name
        if font_size is not None:
            run.font.size = font_size
        color = COLOR_MAP.get(kind)
        if color:
            run.font.color.rgb = color
        if kind == "comment":
            run.font.italic = True
        rpr = run._element.get_or_add_rPr()
        rFonts = rpr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rpr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), font_name)


def is_listing_table(table):
    if len(table.rows) != 1 or len(table.rows[0].cells) != 1:
        return False
    text = table.rows[0].cells[0].text
    return text.strip().startswith("//") and ("right_shifter.v" in text or "PE.v" in text)


def main():
    d = docx.Document(PATH)
    touched = 0
    for table in d.tables:
        if not is_listing_table(table):
            continue
        cell = table.rows[0].cells[0]
        for p in cell.paragraphs:
            colorize_paragraph(p)
        touched += 1
    if touched != 2:
        raise RuntimeError(f"Expected to find exactly 2 listing tables, found {touched}")
    d.save(PATH)
    print(f"Colorized {touched} listing tables in {PATH}")


if __name__ == "__main__":
    main()
