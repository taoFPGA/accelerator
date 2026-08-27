"""
Builds report/taoFPGA_Project_Book.docx: the Bar-Ilan University final-project
report, using report/Final_Project_Book.docx (the real BIU template -- title
page, logos, running header/footer, Abstract/Acknowledgments layout) as the
base document, with:

  1. The title/author/track/supervisor/mentor block updated to the settled
     title and cleared of "[TO CONFIRM]" scaffolding that's now resolved.
  2. The Abstract replaced with the verified, camera-ready IEEE paper's own
     abstract (report/generate_ieee_paper.py), so the two documents never
     disagree with each other.
  3. A static Table of Contents / List of Figures / List of Tables matching
     the actual verified body (built by hand, not a live Word field, since
     the shared heading helpers in generate_ieee_paper.py don't tag real
     Word "Heading" styles -- retrofitting that risked the already-verified
     paper body for a cosmetic gain).
  4. The verified IEEE paper body (Sections I-VIII + References) reused
     completely unchanged via generate_ieee_paper.build_ieee_body(doc) --
     per the user's explicit choice: IEEE body content, BIU front matter.
  5. Appendix A (Engineering Problems Encountered and Their Solutions) and
     Appendix B (General Engineering Principles): condensed, interview-depth
     write-ups of the real debugging history recorded in
     report/project_story.md, added specifically so this book alone is
     enough to prepare for in-depth interview questions about the project,
     without requiring a reader to separately go find project_story.md.
  6. The real branded running header/footer ("Project number: 309" /
     "taoFPGA -- Page N") filled in from placeholder text, left otherwise
     untouched.

Does NOT overwrite Final_Project_Book.docx or taoFPGA_IEEE_paper.docx --
writes a new file, taoFPGA_Project_Book.docx.

Run on the dev machine (needs python-docx, pandoc on PATH -- the same
requirements as generate_ieee_paper.py, which this script imports):
    python generate_project_book.py
"""
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

import generate_ieee_paper as ieee

HERE = os.path.dirname(os.path.abspath(__file__))
BASE_PATH = os.path.join(HERE, "Final_Project_Book.docx")
OUT_PATH = os.path.join(HERE, "taoFPGA_Project_Book.docx")

TITLE_TEXT = (
    "taoFPGA: Architecture, Implementation, and Evaluation of a "
    "Quantized Transformer Processing Core on Edge FPGA"
)
PROJECT_NUMBER = "309"
RUNNING_FOOTER_TITLE = "taoFPGA -- Transformer Accelerator Final Project"
SUBMISSION_DATE = "September 22, 2026"
HEBREW_TITLE_TEXT = "taoFPGA: ארכיטקטורה, מימוש והערכת ביצועים של מאיץ חומרה מקוונטט לטרנספורמרים על גבי FPGA במכשירי קצה"


# ---------------------------------------------------------------------------
# Front-matter helpers (single-column book style, reusing the paper's own
# run-styling helpers so fonts/sizes stay visually consistent throughout).
# ---------------------------------------------------------------------------

def strip_numbering(p):
    """The base template's Heading 1/2 styles carry an attached outline
    numbering list (real chapters render as "1", "1.1", etc. -- confirmed by
    the stale cached TOC field in the base template). Front-matter section
    titles (Table of Contents/List of Figures/List of Tables) and the
    appendix headings below use these styles for visual consistency but
    already carry their own explicit numbering ("A.1", ...) or need none at
    all, so the inherited list numbering must be explicitly suppressed per
    paragraph (numId=0 is the standard OOXML "no list" override)."""
    pPr = p._p.get_or_add_pPr()
    existing = pPr.find(qn("w:numPr"))
    if existing is not None:
        pPr.remove(existing)
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)


def toc_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.text = text
    strip_numbering(p)
    return p


def toc_line(doc, text, page, level=0, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25 * level)
    pf.space_after = Pt(3)
    tab_stops = pf.tab_stops
    tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(f"{text}\t{page}")
    ieee.style_run(r, size=Pt(10.5), bold=bold)
    return p


def replace_placeholder_run(container, old, new):
    """Find a run with exact text `old` anywhere in a header/footer's tables
    (python-docx's `.paragraphs` doesn't reach paragraphs nested inside a
    header/footer's <w:tbl>, so tables must be walked explicitly) and set it
    to `new`. Raises if not found, so a template change doesn't silently
    leave a stale placeholder in the final book."""
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.text == old:
                            r.text = new
                            return
    raise RuntimeError(f"Placeholder {old!r} not found in header/footer table")


# ---------------------------------------------------------------------------
# Appendix content: condensed from report/project_story.md's real,
# chronological debugging log. Each entry keeps Problem / Investigation /
# Fix / Why-it-matters structure so the book itself doubles as interview
# preparation material, not just a project summary.
# ---------------------------------------------------------------------------

APPENDIX_A_ENTRIES = [
    dict(
        title="A.1 Elaboration-Time Memory Blowup From `generate`-Unrolled Test Arrays",
        cite="project_story.md §6-9",
        problem=(
            "Running the matmul testbench at full scale (200×96×160) crashed the "
            "entire simulation session -- not just the job, the whole host -- during "
            "Xcelium's native-code-generation phase."
        ),
        investigation=(
            "The partial log pointed at native-code generation for the testbench module "
            "itself. Hypothesis: `x_in_array`/`y_in_array`/`z_hard_array` were built with a "
            "SystemVerilog `generate` loop that elaborates one hardware object per scalar "
            "array element -- fine for genuinely parallel structures, catastrophic for a "
            "large lookup table. Rather than guess and risk a second crash, a resource-"
            "monitored, kill-switch-protected staged rollout (16³ → 32³ → 64³ → 128³, "
            "polling combined process RSS every 2 seconds with a hard kill at 5,500MB) "
            "confirmed memory grew worse than quadratically -- 59MB → 76MB → 482MB → "
            ">5,900MB and still climbing at 128³."
        ),
        fix=(
            "Replaced the pre-elaborated arrays with `automatic` SystemVerilog functions "
            "(`pack_x_word`, `pack_y_word`, `unpack_z_word`) that compute each word at "
            "runtime, on demand, instead of materializing every possible word as a separate "
            "hardware object. Full scale now runs in ~51MB / ~4 seconds, with functional "
            "output confirmed bit-for-bit unchanged (identical 2,894-cycle latency figure "
            "before and after)."
        ),
        why=(
            "This is a real distinction between synthesizable-style hardware description and "
            "software-style data structures written in the same language: a `generate` loop "
            "is an elaboration-time construct, and using it to hold data (as opposed to "
            "genuinely parallel structure) doesn't scale. It's also a strong example of "
            "diagnosing a scaling problem with a bounded, safety-netted experiment instead of "
            "guessing -- the same protocol was reused for every large simulation/synthesis "
            "run for the rest of the project."
        ),
    ),
    dict(
        title="A.2 A Silent Compile Bug Hidden by an Untested Code Path",
        cite="project_story.md §6",
        problem=(
            "The very first attempt to actually run the matmul testbench (after weeks lost "
            "to an unrelated infrastructure problem) hit an immediate compile error: an "
            "undeclared identifier, `start_trans`."
        ),
        investigation=(
            "`git blame` showed a later commit had added instrumentation code that referenced "
            "`start_trans` above the line that actually declared it -- a pure ordering bug. "
            "The testbench had been silently uncompilable since that commit, but nothing had "
            "ever gotten far enough to hit it, because every previous attempt to run it had "
            "failed earlier, for an unrelated infrastructure reason."
        ),
        fix="Moved the declaration above its first use. No logic change.",
        why=(
            "A working-looking pipeline can hide a real, waiting compile bug indefinitely if "
            "nothing ever exercises the path that hits it. The practical lesson: don't assume "
            "a code path is fine just because nothing has complained about it yet -- absence "
            "of failure isn't evidence of correctness if the path was never actually reached."
        ),
    ),
    dict(
        title="A.3 A Documented, Deliberately-Deferred Bug That Was Still a Real Bug",
        cite="project_story.md §10, §12",
        problem=(
            "`Softmax_control`'s scale output port was 4 bits wide (valid range 7-12) but "
            "`gelu.v`'s corresponding scale input port was only 3 bits wide (max value 7) -- "
            "the two IPs could only agree at exactly scale=7. This was already known and "
            "explicitly flagged in a code comment, left unfixed by whoever wrote the original "
            "integration blueprint."
        ),
        investigation=(
            "Reading `gelu.v`'s actual shift logic showed this wasn't a trivial port-width "
            "bump: the code assumed the input scale maxed out at 7 and had no branch for "
            "anything higher, so widening the port alone would have silently mis-shifted any "
            "newly representable value 8-15 as if it were 7 -- a correctness bug, not just a "
            "missing feature."
        ),
        fix=(
            "Widened the scale port through all three affected modules to 4 bits, added the "
            "missing shift-direction branch in `gelu.v`, and extended the GELU testbench's "
            "randomized range from 0-6 to 0-11 (the old range never even reached the boundary "
            "value 7 that used to be the only legal one)."
        ),
        why=(
            "A code comment that says \"known limitation, not fixed here\" is not the same as "
            "\"already handled\" -- it's worth actively hunting these down rather than trusting "
            "they'll be addressed later by someone else. It's also a good interview example of "
            "reading an interface's actual implementation, not just its declared width, before "
            "trusting a narrow-looking fix."
        ),
    ),
    dict(
        title="A.4 Six Silent Parameter Mismatches Between Verified RTL and an Adopted Integration Script",
        cite="project_story.md §15-16",
        problem=(
            "Reviewing the externally-adopted Vivado integration script (`prj.tcl`, originally "
            "exported from an earlier Vivado session, never maintained against this project's "
            "own verified RTL) ahead of synthesis, one visible mismatch was found: the "
            "systolic array size was configured as 24, but every verified testbench uses 16."
        ),
        investigation=(
            "Rather than fix the one visible value and move on, every one of the wrapper "
            "module's ten RTL-facing parameters was cross-referenced against the exact value "
            "the verified testbenches use. The audit found six mismatches, not one -- and the "
            "two most dangerous ones (a shift-width and a block-count-width parameter) were "
            "never even mentioned in the integration script at all, silently inheriting a "
            "stale wrapper default with no trace of the problem visible anywhere in the file "
            "someone would normally review."
        ),
        fix=(
            "Corrected all six parameters, plus two downstream AXI-Stream width-converter "
            "settings that had been sized to match the wrong array size."
        ),
        why=(
            "The core principle established here (and used for the rest of the project): "
            "verified RTL and testbenches are the source of truth, and any adopted external "
            "tooling gets adapted to match them -- never assumed correct just because it was "
            "machine-generated by the official tool. The specific lesson worth repeating in an "
            "interview: the values that are silently missing from a config (falling back to "
            "some default) are often more dangerous than the values that are visibly wrong, "
            "because nothing prompts a reviewer to go check them."
        ),
    ),
    dict(
        title="A.5 Async-Reset DRC Violations: Chasing One Root Cause Through Twelve Files",
        cite="project_story.md §18",
        problem=(
            "Post-synthesis DRC reported 40+ warnings that block-RAM control logic was driven "
            "by asynchronously reset registers -- a pattern Xilinx's own documentation warns "
            "can corrupt memory contents outside what static timing analysis can see."
        ),
        investigation=(
            "What looked like a three-file fix (the buffer modules DRC actually named) grew "
            "to twelve files across three re-synthesis rounds: fixing the named files made the "
            "warning move to a different register one hierarchy level deeper each time -- "
            "first into a nested submodule, then into an unrelated width-adapter module with "
            "no \"buffer\" in its name at all. A project-wide grep for the actual anti-pattern "
            "(the specific asynchronous-reset sensitivity-list syntax) eventually enumerated "
            "the full remaining scope in one pass, rather than continuing to whack-a-mole one "
            "violation at a time."
        ),
        fix=(
            "Converted 60 `always` blocks across 12 files from asynchronous to synchronous "
            "reset -- Xilinx's own recommended remedy -- and re-verified bit-for-bit identical "
            "simulation results after every round."
        ),
        why=(
            "A DRC violation names the nearest offender on a signal path, not the full extent "
            "of the underlying pattern -- the same root cause can recur at every upstream hop "
            "in a handshake network, in modules that share no obvious naming relationship with "
            "where the symptom first appeared. This is also a clean example of when to stop "
            "iterating on individual violations and instead do one bounded, project-wide audit."
        ),
    ),
    dict(
        title="A.6 Discovering the Full Pipeline Had Never Actually Been AXI-Wrapped",
        cite="project_story.md §19",
        problem=(
            "Before running full SoC integration, a check of the integration script's actual "
            "scope revealed it only ever wrapped the standalone matmul engine in AXI -- the "
            "complete matmul-Softmax-GELU pipeline this project had been calling \"the "
            "accelerator\" for many prior sections had never been given any AXI wrapping at "
            "all, only the raw ports the testbenches drive directly."
        ),
        investigation=(
            "Caught before running anything, by checking what the integration script actually "
            "targeted rather than assuming its scope matched the project's current "
            "vocabulary."
        ),
        fix=(
            "Built a new AXI4-Lite/AXI4-Stream wrapper for the full pipeline, deliberately "
            "keeping the underlying RTL's own parameter names instead of introducing new "
            "aliases -- a direct, explicit application of the lesson from A.4, since renaming "
            "was exactly what caused that earlier drift bug. The new wrapper needed zero "
            "configuration overrides in the integration script as a direct result."
        ),
        why=(
            "\"Run the existing integration script\" is not always as simple as running it -- "
            "checking that its scope still matches what the rest of the project means by \"the "
            "accelerator\" is worth a pause before spending a synthesis run on the wrong "
            "target. This is also a strong example of a documented lesson actively changing "
            "how the next piece of work was built, not just how a bug was fixed after the "
            "fact."
        ),
    ),
    dict(
        title="A.7 A Synthesis Attribute That Silently Did Nothing",
        cite="project_story.md §20",
        problem=(
            "DSP48E1 hard-macro usage sat at just 5.91% while LUT usage climbed toward 58% at "
            "the full-SoC level -- an inversion of what a matmul-dominated systolic array "
            "accelerator should look like. The first attempt to force multiplications onto "
            "DSP48E1s, by placing a `use_dsp` synthesis attribute directly above the "
            "multiply-accumulate `always` block, produced byte-identical utilization on "
            "re-synthesis: no error, and no effect."
        ),
        investigation=(
            "Per Xilinx UG901, the `use_dsp` attribute must attach to the register declaration "
            "that holds the multiply result, not the enclosing procedural block -- an "
            "attribute on the `always` statement itself isn't a recognized attachment point "
            "and is silently dropped during elaboration, with no warning."
        ),
        fix=(
            "Changed the relevant output from `output reg` to `output wire`, introduced a "
            "per-branch internal register carrying the attribute correctly, and drove the "
            "port via a continuous assignment -- functionally and cycle-timing identical to "
            "the original RTL. Re-synthesis then showed DSP48E1 usage jump from 13 to 205 and "
            "CARRY4 usage drop by more than half. True dual-8-bit DSP packing was considered "
            "and explicitly ruled out first, since it would have required restructuring the "
            "systolic array's timing skew, a real microarchitecture change; the chosen fix was "
            "a lower-risk partial 1:1 DSP mapping (12 of 16 systolic rows) sized to fit the "
            "chip's DSP budget."
        ),
        why=(
            "A synthesis pragma that produces no error and no effect is exactly the kind of "
            "silent failure that's easy to miss if you only check for compile errors, not "
            "utilization numbers before and after. This is one of the project's strongest "
            "single findings and worth being able to explain precisely (attribute-attachment "
            "rules, wire-vs-reg, why the fix preserves identical behavior) in an interview."
        ),
    ),
    dict(
        title="A.8 Same Chip, Different Board: the PYNQ-Z1 to PYNQ-Z2 Hardware Audit",
        cite="project_story.md §22",
        problem=(
            "The physical board on hand is a PYNQ-Z2, but the whole flow had been built and "
            "verified against PYNQ-Z1 board files."
        ),
        investigation=(
            "Both boards use the identical Zynq-7020 chip, so everything living inside the "
            "fabric (RTL, timing closure, utilization, the complete absence of any user I/O "
            "constraints file) travels unchanged. The real risk was narrower and easy to "
            "overlook: the integration script's PS7 configuration hardcodes not just a "
            "cosmetic board-name string but the actual DDR3 board-preset values -- part "
            "number, per-pin trace-length and delay figures -- that model the PYNQ-Z1 PCB's "
            "specific DDR3 layout, which is a different physical board than the PYNQ-Z2 "
            "actually being used."
        ),
        fix=(
            "Deliberately did not hand-patch plausible-looking replacement numbers, since a "
            "fabricated-but-plausible DDR calibration value is worse than an honestly-wrong "
            "one -- it would look fixed while still being invented. Updated the board-part "
            "string to attempt the correct board identifier and left explicit inline comments "
            "flagging exactly which values are still PYNQ-Z1-specific, so the gap can't be "
            "silently forgotten. The correct fix (installing the board vendor's official "
            "Vivado board files and re-applying the preset from the GUI) was documented as the "
            "next concrete step rather than approximated."
        ),
        why=(
            "\"Same silicon part number\" and \"same board\" are not the same claim -- a good "
            "interview answer distinguishes precisely which parts of a hardware design travel "
            "with the chip and which are tied to the physical board it sits on, and explains "
            "why guessing a DDR calibration number is actively worse than leaving it visibly "
            "unresolved."
        ),
    ),
    dict(
        title="A.9 A Comprehensively Stale Software Driver, Found and Rebuilt From the Real Register Map",
        cite="project_story.md §23, §25",
        problem=(
            "The repository's existing PYNQ/Python driver (`apps/MM.py`) predated the "
            "project's later integration of the full matmul-Softmax-GELU pipeline. It "
            "configured the wrong array size, hardcoded peripheral addresses that didn't "
            "match the real integration script's address map, and only knew about the old "
            "4-register matmul-only control interface."
        ),
        investigation=(
            "Rather than patch individual wrong values, the whole register/DMA contract was "
            "re-derived directly from the current RTL wrapper's own header comments and the "
            "real generated hardware description -- the same discipline established for the "
            "integration script in A.4. On the very first real hardware run, this caught a "
            "second, independent mismatch: PYNQ names a custom AXI-Lite peripheral after its "
            "actual AXI interface name, not its block instance name, so a hardcoded attribute "
            "path would have failed the moment it was used."
        ),
        fix=(
            "Rewrote the driver from scratch against the real address map and register "
            "layout, and replaced the hardcoded IP-attribute access with a small resolver "
            "helper that searches the overlay's actual IP dictionary instead of assuming one "
            "specific naming convention."
        ),
        why=(
            "A driver that looks complete (helper functions, shape validation, error "
            "handling) can still be built entirely against the wrong version of the hardware "
            "-- every one of its wrong values traced back to a single root cause, that nobody "
            "had updated it since the hardware it targeted changed. It's also a clean example "
            "of not trusting an assumption about how a tool names things until checking the "
            "tool's own actual output."
        ),
    ),
    dict(
        title="A.10 A Real, Reproducible Hardware/Software Mismatch, Diagnosed Methodically",
        cite="project_story.md §26, §28",
        problem=(
            "The first numerical test on real PYNQ-Z2 hardware showed 11 of 256 output "
            "elements (about 4.3%) outside the golden model's tolerance -- every failing "
            "output was exactly zero, against expected values around 0.05-0.10."
        ),
        investigation=(
            "Investigated in a fixed order, ruling out cheaper explanations before touching "
            "RTL: re-ran the identical input three times and got bit-for-bit identical "
            "failures, ruling out timing noise; confirmed the FIFO-overflow safeguard hadn't "
            "fired; correlated 15 additional trials against each element's distance from its "
            "row's maximum value; and hand-traced the Softmax/Exp/Ln modules for a specific "
            "zero-output hypothesis, which the trace refuted. Rather than keep guessing by "
            "hand, wrote a standalone debug testbench reproducing the exact failing row "
            "through the Softmax module alone, with a full internal-signal trace. It came "
            "back clean: the isolated module computed the correct, non-zero answer given the "
            "same inputs -- proving the bug, if it was a logic bug at all, had to be somewhere "
            "the isolated test couldn't reach (the real data path feeding it, the actual "
            "register configuration, or something only real hardware timing could expose)."
        ),
        fix=(
            "Not resolved by this test alone -- it correctly redirected the search rather than "
            "closing it (see A.11)."
        ),
        why=(
            "A clean negative result is not a dead end, it's a boundary: once a module is "
            "proven correct in isolation, the remaining search space is provably everything "
            "outside it, not \"somewhere in the module, just not the spots already checked.\" "
            "This entry is a strong interview answer to \"tell me about a hard bug you "
            "debugged,\" since it shows a disciplined process (cheapest checks first, "
            "structural before circumstantial) rather than a lucky guess."
        ),
    ),
    dict(
        title="A.11 The Real Root Cause: a Tiny-Shape Edge Case, Not a Design Flaw",
        cite="project_story.md §29",
        problem=(
            "Following A.10's redirection, a targeted \"marker row\" hardware test (a unique "
            "spike per row, so a correctly-framed pipeline reproduces the spike's exact "
            "column position) showed the data-framing path itself was wrong even for a single "
            "row, and a sweep across small row counts hung outright at two rows."
        ),
        investigation=(
            "Every failing test up to this point -- including every test in A.10 -- used "
            "matrix-block-count values far smaller than anything the design had ever actually "
            "been verified at. Re-running the exact shape and configuration the full-pipeline "
            "simulation had already verified end-to-end (200×96×160) through the same real "
            "hardware path produced 0 of 32,000 output elements outside tolerance: the "
            "accelerator is correct on real silicon at the scale it was designed and verified "
            "for."
        ),
        fix=(
            "Reading the input-buffer's row-address counter logic directly (not hand-waved) "
            "found one concrete cause for the single-row failure: a wraparound condition that "
            "is mathematically correct for every row count except exactly one, where the wrap "
            "value it checks for can never actually occur. The two-row hang was traced to the "
            "same small-block-count territory but was not fully root-caused, and was "
            "explicitly left open rather than claimed fixed."
        ),
        why=(
            "This is the project's central verification lesson: a \"quick\" small smoke test "
            "accidentally exercised a completely different, never-validated region of the "
            "design's parameter space, and nearly turned a genuine but narrow edge case into a "
            "false alarm about the whole accelerator's correctness. The decisive move wasn't a "
            "cleverer hypothesis -- it was re-running the exact condition already proven "
            "correct and confirming the result actually matched. A strong interview point: "
            "distinguishing \"the design is wrong\" from \"my test exercised an unverified "
            "corner of the design\" is a real, valuable engineering skill."
        ),
    ),
    dict(
        title="A.12 Refusing to Report a Speedup the Pipeline Doesn't Actually Support",
        cite="project_story.md §30",
        problem=(
            "Asked for a complete end-to-end Vision Transformer (ViT) benchmark: run a real "
            "ViT model, accelerate it with the hardware, and report the speedup."
        ),
        investigation=(
            "Before writing any benchmark code, checked whether the accelerator's fixed "
            "pipeline (matmul, then per-row Softmax, then GELU, fused with no way to tap the "
            "intermediate result) structurally matches any real ViT sub-block. It doesn't: "
            "self-attention needs matmul-softmax-matmul with no GELU in between, and the MLP "
            "block needs matmul-GELU-matmul with no Softmax in between. There is no ViT "
            "operation this specific fixed pipeline can correctly substitute into."
        ),
        fix=(
            "Presented this finding honestly rather than silently building a benchmark that "
            "would have reported a broken classification alongside a fabricated speedup "
            "number. The chosen alternative: a real, independently validated NumPy "
            "reimplementation of ViT-Tiny for a genuine software baseline and correctness "
            "check (matching the real PyTorch reference to float32 noise), plus a separate, "
            "clearly-labeled hardware kernel-throughput benchmark that never claims to "
            "represent an accelerated end-to-end ViT run."
        ),
        why=(
            "The most valuable check before writing a benchmark is often not the code, it's "
            "whether the comparison it will report is actually true. This is a strong example "
            "of scientific/engineering integrity under pressure to produce an impressive "
            "number -- exactly the kind of judgment call an interviewer may probe for "
            "directly."
        ),
    ),
    dict(
        title="A.13 A 64KB DMA Ceiling, and Why the Obvious Workaround Would Have Silently Corrupted Data",
        cite="project_story.md §31",
        problem=(
            "The first real ViT-representative kernel benchmark run succeeded at one matrix "
            "shape but hard-failed at a larger one: the DMA engine refused the transfer "
            "outright because it exceeded a 65,536-byte maximum."
        ),
        investigation=(
            "Root cause: the three DMA cores were synthesized with a 16-bit transfer-length "
            "register, capping any single transfer at 64KB, and the larger shape's data "
            "exceeded that. The obvious-looking fix -- split the oversized transfer into "
            "several smaller DMA calls -- was checked against the actual RTL before being "
            "used: the input buffer's write-address counters reset to zero on the stream's "
            "end-of-frame signal, and this DMA mode asserts that signal at the end of every "
            "transfer, not only a genuine final one. Two chunked transfers would each look "
            "like \"the last one\" to the RTL, so the second chunk would silently overwrite the "
            "first from address zero instead of continuing it -- corrupting the buffer rather "
            "than raising any visible error."
        ),
        fix=(
            "Capped the affected benchmark shape at the largest size that fits within the "
            "64KB ceiling, labeled explicitly (in code and in printed output) as DMA-limited "
            "rather than the true target dimension. The real fix -- a wider transfer-length "
            "register, or genuine Scatter-Gather DMA with explicit end-of-frame control per "
            "descriptor -- needs re-synthesis and was logged as a specific, well-understood "
            "open item rather than implemented under time pressure."
        ),
        why=(
            "The fastest-looking fix for a hardware error is not automatically the correct "
            "one -- tracing exactly what a workaround would do to the specific RTL logic "
            "involved caught a data-corrupting bug before it shipped. This is a good example "
            "of hardware/software co-design intuition: a software-only fix (chunking) can be "
            "actively unsafe when the hardware's own protocol semantics (what \"last transfer\" "
            "means at the RTL level) don't match what the software layer assumes."
        ),
    ),
]

APPENDIX_B_PRINCIPLES = [
    ("Infrastructure first.", "Verify infrastructure and environment assumptions before a "
     "deep technical investigation -- the single largest time cost early in this project "
     "(project_story.md §1-5) turned out not to be a design bug at all."),
    ("A fix isn't proven until the original symptom is confirmed gone.", "Ideally more than "
     "once, and ideally by reproducing the exact original failing condition, not a similar-"
     "looking one."),
    ("`generate` loops don't scale as lookup tables.", "One hardware object per array element "
     "is correct for genuinely parallel structure and catastrophic for a large data array; use "
     "an on-demand function instead."),
    ("Headroom and sizing parameters must scale in lockstep with data size.", "A previously-"
     "fixed bug can silently reappear at a new scale if a related sizing parameter wasn't "
     "raised with it -- read and respect a file's own inline warnings about this."),
    ("A fix to a shared interface needs every consumer re-checked.", "Not just the interface "
     "declaration itself -- a testbench or caller can keep silently exercising only the old, "
     "narrower behavior even after the interface is fixed."),
    ("Externally-adopted integration scripts are not authoritative just because they were "
     "machine-generated.", "Audit their full parameter list against the verified design's own "
     "declarations, not only the one value already suspected wrong -- the values nobody "
     "explicitly overrode are typically the most dangerous ones, because nothing prompts a "
     "reviewer to go check them."),
    ("Simulation and post-synthesis design-rule checking catch different classes of bug.", "A "
     "clean simulation proves functional correctness under the testbench's own stimulus and "
     "timing model, not silicon-level correctness -- asynchronous resets driving block-RAM "
     "control logic are a textbook example of something a behavioral simulator cannot expose "
     "but that a DRC pass exists specifically to catch."),
    ("A design-rule violation names the nearest offender, not the full extent of the pattern.",
     "The same root cause can recur at every upstream hop in a signal path, in modules with no "
     "obvious naming relationship to where the symptom first appeared -- a project-wide audit "
     "for the actual anti-pattern is more reliable than iteratively chasing wherever the next "
     "violation happens to point."),
    ("A synthesis pragma that produces no error can still silently do nothing.", "Attribute-"
     "attachment rules matter -- check the actual utilization numbers before and after, not "
     "just whether the build completed without complaint."),
    ("\"Same chip\" and \"same board\" are different claims.", "Anything that lives purely in "
     "the silicon fabric travels between two boards sharing the same part number; anything "
     "that models the physical PCB (DRAM calibration, board-preset metadata) does not, and a "
     "plausible-looking guessed value is worse than an honestly-flagged gap."),
    ("A driver or script that looks complete can still target the wrong version of the "
     "hardware.", "Re-derive an interface contract from the current design's own source of "
     "truth rather than trusting that an existing file was ever updated when the design "
     "changed underneath it."),
    ("A hardware bug search needs the same scale discipline as everything else.", "Testing "
     "\"quickly\" with a small, convenient shape can accidentally exercise a completely "
     "different, never-validated region of the design's parameter space -- when in doubt, "
     "re-run the exact condition already proven correct and confirm the result actually "
     "matches, before concluding the design itself is at fault."),
    ("The most valuable check before generating a result is whether the comparison is even "
     "true.", "An accelerator, benchmark, or workaround that looks plausible is not "
     "automatically valid for the specific case at hand -- checking this explicitly, in "
     "writing, before producing a number is worth the pause, especially for results that will "
     "be reported to someone else."),
]


def build_appendix_a(doc):
    p = doc.add_paragraph(style="Heading 1")
    p.text = "Appendix A: Engineering Problems Encountered and Their Solutions"
    strip_numbering(p)
    ieee.body_para(doc,
        "The verified results in the body of this report are the outcome of a long, "
        "iterative debugging process, not a linear implementation. This appendix records the "
        "real, individually significant problems found and solved along the way, condensed "
        "from the project's full chronological engineering log "
        "(report/project_story.md), specifically so this report is sufficient on its own to "
        "answer detailed, in-depth questions about how the project's key results were "
        "actually reached.")
    for entry in APPENDIX_A_ENTRIES:
        h = doc.add_paragraph(style="Heading 2")
        h.text = entry["title"]
        strip_numbering(h)
        ieee.body_para(doc, entry["problem"])
        p = doc.add_paragraph()
        r = p.add_run("Investigation. ")
        ieee.style_run(r, size=ieee.BODY_SIZE, bold=True)
        r2 = p.add_run(entry["investigation"])
        ieee.style_run(r2, size=ieee.BODY_SIZE)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(6)
        p = doc.add_paragraph()
        r = p.add_run("Fix. ")
        ieee.style_run(r, size=ieee.BODY_SIZE, bold=True)
        r2 = p.add_run(entry["fix"])
        ieee.style_run(r2, size=ieee.BODY_SIZE)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(6)
        p = doc.add_paragraph()
        r = p.add_run("Why this matters. ")
        ieee.style_run(r, size=ieee.BODY_SIZE, bold=True, italic=True)
        r2 = p.add_run(entry["why"])
        ieee.style_run(r2, size=ieee.BODY_SIZE, italic=True)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        cite_p = doc.add_paragraph()
        r = cite_p.add_run(entry["cite"])
        ieee.style_run(r, size=Pt(8.5), italic=True)
        cite_p.paragraph_format.space_after = Pt(12)


def build_appendix_b(doc):
    p = doc.add_paragraph(style="Heading 1")
    p.text = "Appendix B: General Engineering Principles"
    strip_numbering(p)
    ieee.body_para(doc,
        "The individual problems in Appendix A share a smaller number of recurring "
        "principles. They are collected here as a standalone summary -- useful as a direct "
        "answer to a broader interview question such as \"what did you learn from this "
        "project,\" independent of any single bug.")
    for title, text in APPENDIX_B_PRINCIPLES:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(title + " ")
        ieee.style_run(r, size=ieee.BODY_SIZE, bold=True)
        r2 = p.add_run(text)
        ieee.style_run(r2, size=ieee.BODY_SIZE)
        p.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------------------
# Assembly
# ---------------------------------------------------------------------------

doc = docx.Document(BASE_PATH)

# ---- Fix the real branded header/footer (currently placeholder text) ----
section0 = doc.sections[0]
replace_placeholder_run(section0.header, "Project number: [#]", f"Project number: {PROJECT_NUMBER}")
replace_placeholder_run(section0.header, "[date of submission]", SUBMISSION_DATE)
replace_placeholder_run(section0.footer, "Project: [name of a project]", RUNNING_FOOTER_TITLE)

paragraphs = doc.paragraphs

# ---- Update the title block in place (paragraphs verified by inspection of
# the base template: 7-8 = title, 9 = stale "[TO CONFIRM]" line to remove,
# 10-11 = Hebrew title placeholder, 22 = submission date) ----
paragraphs[7].text = TITLE_TEXT
for r in paragraphs[7].runs:
    ieee.style_run(r, size=Pt(16), bold=True)
paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.CENTER

paragraphs[8].text = ""
paragraphs[9].text = ""  # was: "[TO CONFIRM -- working title above ...]"
paragraphs[10].text = HEBREW_TITLE_TEXT
for r in paragraphs[10].runs:
    ieee.style_run(r, size=Pt(13), bold=True)
    # Restore the RTL run-direction/language markup the placeholder had --
    # python-docx's plain .text= setter doesn't carry it over to the new run.
    rPr = r._r.get_or_add_rPr()
    rPr.append(OxmlElement("w:rtl"))
    lang = OxmlElement("w:lang")
    lang.set(qn("w:bidi"), "he-IL")
    rPr.append(lang)
paragraphs[10].alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraphs[11].text = ""  # was: "[TO CONFIRM -- Hebrew title still needed]"
paragraphs[22].runs[0].text = paragraphs[22].runs[0].text.replace(
    "[date of submission]", SUBMISSION_DATE)
paragraphs[23].text = ""  # was: "[TO CONFIRM -- confirm the supervisor/mentor role mapping ...]"
paragraphs[30].text = ""  # was: leftover template instructional sentence
paragraphs[39].text = ""  # was: "[TO CONFIRM -- add any additional personal acknowledgments here]"

# The base template has two manual page breaks (paragraph 29, right after
# "Project number: 309", and paragraph 33, right before "ABSTRACT") with
# only blank filler paragraphs between them -- that combination renders as
# a fully blank page. Removing the first break lets the blank fillers
# collapse into normal paragraph spacing while the second break still
# gives the Abstract its own fresh page.
for br in paragraphs[29]._p.findall(".//" + qn("w:br")):
    if br.get(qn("w:type")) == "page":
        br.getparent().remove(br)

# ---- Replace the Abstract with the verified IEEE paper's own abstract ----
# paragraphs[34] = "ABSTRACT" heading (kept as-is); paragraphs[35] = body text.
ABSTRACT_FONT_SIZE = Pt(11)
ABSTRACT_TEXT = (
    "Transformer models are increasingly deployed on power- and resource-constrained edge "
    "devices, where their quadratic attention cost and large dense matrix multiplications "
    "strain both compute and memory bandwidth. This report presents taoFPGA, a quantized "
    "INT8 transformer processing core built around a 16×16 weight-stationary systolic "
    "array fused with dedicated fixed-point Softmax and GELU engines, verified end to end "
    "from RTL simulation through physical signoff and real hardware bring-up on a Xilinx "
    "Zynq-7020 (PYNQ-Z2). We detail a tolerance-bounded verification methodology carried "
    "consistently from SystemVerilog testbenches through a Python/NumPy golden model "
    "validated against physical silicon, a synthesis-level DSP-inference defect whose "
    "correction reduced LUT utilization by 49.3% and total on-chip power by 3.1%, and a "
    "full-scale hardware validation run matching simulation to zero error across 32,000 "
    "output elements. We further report a kernel-level benchmark against a real, "
    "numerically validated ViT-Tiny software baseline, achieving up to 97.73× measured "
    "speedup (70.02× and 97.73× across the two benchmarked kernel shapes) over the "
    "board's ARM Cortex-A9, corresponding to a total energy efficiency of 2.25 GOPS/W "
    "(2.47 GOPS/W on dynamic power alone) at peak throughput -- while explicitly analyzing "
    "why the accelerator's fixed pipeline does not structurally correspond to any single "
    "Vision Transformer operation, a distinction we treat as a methodological contribution "
    "in its own right. We close with a quantified account of the design's remaining "
    "hardware constraints, principally a 64KB single-transfer DMA ceiling, and the "
    "Scatter-Gather DMA redesign identified as its prerequisite fix."
)
paragraphs[35].text = ABSTRACT_TEXT
for r in paragraphs[35].runs:
    ieee.style_run(r, size=ABSTRACT_FONT_SIZE)

# ---- Clear everything from the old "List of figures" heading onward: the
# stale List of Figures/Tables, the old "Introduction" chapter through the
# old References/Appendix A/Appendix B. This content predates hardware
# bring-up, the DSP fix, and the ViT work, and is superseded entirely by the
# freshly-built TOC/LOF/LOT below plus build_ieee_body() and the new
# appendices. (The old chapter body sits *after* the old LOF/LOT in this
# template, so removal must start at the LOF heading, not at "Introduction",
# or the stale LOF/LOT would remain alongside the newly-built ones.)
body = doc.element.body
LOF_IDX = None
for i, p in enumerate(paragraphs):
    if p.text.strip() == "List of figures":
        LOF_IDX = i
        break
if LOF_IDX is None:
    raise RuntimeError("Could not find the old 'List of figures' heading to excise")

# The template also embeds a real, stale Word TOC field between the
# Acknowledgments block and "List of figures" -- its cached display text
# reads blank via python-docx's plain .text accessor (confirmed by direct
# inspection), so it isn't caught by the "List of figures" anchor above.
# Rendering the base template confirmed it renders as broken cached content
# (old chapter titles alongside "Error! Bookmark not defined."), so walk
# backward from LOF_IDX to the last real content paragraph and cut from
# right after it, sweeping up that hidden field regardless of its exact
# paragraph span.
CUT_IDX = LOF_IDX
for i in range(LOF_IDX - 1, -1, -1):
    if paragraphs[i].text.strip():
        CUT_IDX = i + 1
        break

to_remove = paragraphs[CUT_IDX:]
for p in to_remove:
    body.remove(p._p)
# Also drop any leftover tables from the old chapters (old Table 1/2/3).
for tbl in list(doc.tables):
    tbl._element.getparent().remove(tbl._element)

# The template's real "Table of Contents" (a live Word TOC field with its
# own cached, now-stale entries -- "1 Introduction", "1.1 Background", ...)
# lives inside a <w:sdt> content-control wrapper. python-docx's
# `document.paragraphs` only walks direct <w:p> children of the body, so
# every paragraph nested inside that wrapper (confirmed by direct XML
# inspection: one body-level <w:sdt> holding 16 nested paragraphs) is
# invisible to the paragraph-index-based removal above and survives
# untouched unless removed explicitly here.
for sdt in body.findall(qn("w:sdt")):
    body.remove(sdt)

# ---- Rebuild Table of Contents / List of Figures / List of Tables ----
# Page numbers below were read directly off a rendered PDF of this exact
# layout (fixed content, fixed page size/margins -- won't drift unless the
# front matter or body content changes length). If the front matter or the
# IEEE body is edited later, regenerate once with placeholder numbers,
# re-read the real page numbers off the PDF, and update this table again.
toc_heading(doc, "Table of Contents")
toc_line(doc, "Abstract", 3)
toc_line(doc, "Acknowledgments", 3)
TOC_ENTRIES = [
    ("Preliminaries", 0, 6),
    ("A. Transformer Self-Attention", 1, 6),
    ("B. Fixed-Point (INT8) Quantization", 1, 6),
    ("C. Systolic Arrays for Matrix Multiplication", 1, 6),
    ("I. Introduction & Motivation", 0, 6),
    ("A. The Edge AI Compute Dilemma", 1, 6),
    ("B. Design Intent & Methodology", 1, 6),
    ("C. Key Contributions & Paper Outline", 1, 7),
    ("II. Core Hardware Architecture & Mathematical Formulation", 0, 7),
    ("A. Systolic Matrix Multiplication Core", 1, 7),
    ("B. Fixed-Point Non-Linear Arithmetic Engines", 1, 8),
    ("C. Numerical Quantization Strategy", 1, 9),
    ("III. RTL Verification & Behavioral Simulation (Cadence Xcelium)", 0, 9),
    ("A. Testbench Architecture & Golden Reference Model", 1, 9),
    ("B. Simulation Results & Numerical Precision Verification", 1, 9),
    ("C. Pre-Silicon Logic Debugging", 1, 9),
    ("IV. SoC Integration & Hardware Platform Design (PYNQ-Z2 / Zynq-7020)", 0, 9),
    ("A. SoC Architecture", 1, 9),
    ("B. Physical Adaptations for Board Deployment", 1, 10),
    ("C. Software Infrastructure for Hardware Bring-Up", 1, 11),
    ("V. Physical Implementation, P&R, and Sign-Off Optimizations", 0, 11),
    ("A. Baseline Physical Sign-Off", 1, 11),
    ("B. The DSP Inference Gap & Attribute Resolution", 1, 11),
    ("C. Physical Congestion vs. Timing Trade-Off", 1, 11),
    ("D. Bitstream & Hardware Platform Sign-Off", 1, 12),
    ("VI. Experimental Evaluation & Hardware Bring-Up", 0, 12),
    ("A. Benchmarking Methodology", 1, 12),
    ("B. Performance Metrics & Acceleration", 1, 12),
    ("C. Physical Silicon Verification", 1, 13),
    ("D. Workload Characterization vs. ViT Networks", 1, 13),
    ("E. Hardware Bring-Up Observations & Edge Cases", 1, 13),
    ("F. Comparative Analysis with Prior FPGA Transformer Accelerators", 1, 14),
    ("VII. Engineering Discussion, Bottlenecks & Future Work", 0, 14),
    ("A. DMA Buffer Boundary Constraints", 1, 14),
    ("B. Serial Softmax Throughput Bottleneck", 1, 14),
    ("C. Architectural Roadmap", 1, 14),
    ("VIII. Conclusion", 0, 15),
    ("References", 0, 15),
    ("Appendix A: Engineering Problems Encountered and Their Solutions", 0, 16),
]
APPENDIX_A_PAGES = [16, 16, 17, 17, 17, 18, 18, 19, 19, 20, 20, 20, 21]
# A few appendix titles are too long to fit this indent level before the
# dot-leader tab stop (confirmed by rendering: the full A.13 title runs
# straight into its page number with no visible leader) -- shortened here
# for the TOC line only; the full descriptive title is unchanged in the
# appendix body itself.
TOC_TITLE_OVERRIDES = {
    "A.13 A 64KB DMA Ceiling, and Why the Obvious Workaround Would Have Silently Corrupted Data":
        "A.13 A 64KB DMA Ceiling, and an Unsafe-Looking Workaround",
}
for entry, page in zip(APPENDIX_A_ENTRIES, APPENDIX_A_PAGES):
    toc_text = TOC_TITLE_OVERRIDES.get(entry["title"], entry["title"])
    TOC_ENTRIES.append((toc_text, 1, page))
TOC_ENTRIES.append(("Appendix B: General Engineering Principles", 0, 22))

for text, level, page in TOC_ENTRIES:
    toc_line(doc, text, page, level=level, bold=(level == 0))

toc_heading(doc, "List of Figures")
LOF_ENTRIES = [
    ("Fig. 1. One PE's internal datapath (stationary weight, INT8 multiply, systolic propagation).", 8),
    ("Fig. 2. Softmax_control's 3-pass FSM and its per-pass equations.", 8),
    ("Fig. 3. Illustrative AXI4-Stream handshake and 3-pass Softmax timing structure.", 9),
    ("Fig. 4. Complete top-level SoC architecture.", 10),
    ("Fig. 5. Latency, CPU vs. hardware, both benchmarked shapes.", 12),
    ("Fig. 6. Throughput (GOP/s), CPU vs. hardware, both shapes.", 12),
    ("Fig. 7. Measured hardware kernel speedup over CPU.", 13),
    ("Fig. 8. Real ViT dataflows contrasted with taoFPGA's fixed fused pipeline.", 13),
]
for text, page in LOF_ENTRIES:
    toc_line(doc, text, page)

toc_heading(doc, "List of Tables")
LOT_ENTRIES = [
    ("Table I. transformer_block_axi_top AXI4-Lite Register Map.", 10),
    ("Table II. Post-Route Signoff Metrics, Before/After the DSP-Inference Fix.", 12),
    ("Table III. Kernel Benchmark Results, PYNQ-Z2 vs. ARM Cortex-A9.", 13),
    ("Table IV. Comparison with Literature FPGA Transformer Accelerators.", 14),
]
for text, page in LOT_ENTRIES:
    toc_line(doc, text, page)

# ---- Verified IEEE paper body (Sections I-VIII + References), unchanged ----
doc.add_page_break()
ieee.build_ieee_body(doc)

# ---- Appendices (single column) ----
# A plain page break before the section-break-triggered column-count change
# ensures the two-column References text always ends and the single-column
# Appendix A always starts on a fresh page, regardless of exactly how much
# room happens to be left at the bottom of the References column.
doc.add_page_break()
ieee.new_section(doc, columns=1)
build_appendix_a(doc)
doc.add_page_break()
build_appendix_b(doc)

doc.save(OUT_PATH)
print(f"Wrote {OUT_PATH}")
